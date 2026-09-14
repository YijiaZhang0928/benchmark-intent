#!/usr/bin/env python3
"""Freeze the first five released exact-pair rows, excluding the observed T09 case."""

from __future__ import annotations

import hashlib
import json
import re
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
WORKBOOK = Path("/Users/lora/Documents/Codex/2026-09-12/a/outputs/0913_pdr_bench_15_cold_start_and_micro_rubrics_v0_82.xlsx")
PDR = PROJECT / "tmp/pdr-bench/data"
SELECTED = ["T01", "T02", "T05", "T08", "T11"]
ASKABLE = {
    "T01": {"T1-P1", "T1-P2", "T1-P3", "T1-P4"},
    "T02": {"T2-P1", "T2-P2", "T2-P3", "T2-P4", "T2-P5"},
    "T05": {"T5-P1", "T5-P2", "T5-P3", "T5-P5"},
    "T08": {"T8-P1", "T8-P2", "T8-P3", "T8-P4", "T8-P5"},
    "T11": {"T11-P1", "T11-P2", "T11-P3", "T11-P5"},
}


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def parse_units(text: str, impact: str, askable: set[str]) -> list[dict]:
    units = []
    for line in str(text).splitlines():
        match = re.match(r"(T\d+-(?:P|A)\d+)\s+(.+)", line.strip())
        if match:
            units.append({"id": match.group(1), "impact": impact, "askable_high": match.group(1) in askable, "label": match.group(2)})
    return units


def main() -> int:
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=False)
    task_rows = list(workbook["15 Tasks"].iter_rows(values_only=True))
    task_header = list(task_rows[0])
    rubric_rows = list(workbook["Detailed Rubrics"].iter_rows(values_only=True))
    rubric_header = list(rubric_rows[3])
    personas = read_jsonl(PDR / "persona_data/personas_en.jsonl")
    criteria = read_jsonl(PDR / "criteria_data/criteria150_en.jsonl")
    manifest = []
    for row_number, row in enumerate(task_rows[1:], 2):
        record = dict(zip(task_header, row))
        match = re.search(r"Benchmark Task (\d+) \| PDR source Task (\d+) \| (User\d+)", str(row[0]))
        if not match:
            continue
        task_key = f"T{int(match.group(1)):02d}"
        if task_key not in SELECTED:
            continue
        pdr_task = int(match.group(2)); user = match.group(3)
        case_root = ROOT / "cases" / task_key
        task_dir = case_root / "task"
        task_dir.mkdir(parents=True, exist_ok=True)
        strict = [dict(zip(rubric_header, item)) for item in rubric_rows[4:] if item and item[0] == task_key and item[3] == "P_strict"]
        official_rows = [dict(zip(rubric_header, item)) for item in rubric_rows[4:] if item and item[0] == task_key and item[3] == "P_official"]
        if len(strict) != 67 or not official_rows:
            raise RuntimeError(f"Unexpected rubrics for {task_key}: {len(strict)}/{len(official_rows)}")
        persona = next(item for item in personas if item.get("userid") == user)
        official = next(item for item in criteria if item.get("taskid") == pdr_task and item.get("userid") == user)
        instruction = str(record["input/task_instruction for cold start"]).strip()
        units = parse_units(record["GT_high_impact_preferences"], "high", ASKABLE[task_key]) + parse_units(record["average_impact_preferences"], "average", ASKABLE[task_key])
        if len(units) != 8:
            raise RuntimeError(f"Expected eight units for {task_key}, got {len(units)}")
        files = {
            "instruction.txt": instruction + "\n",
            "hidden_persona.json": json.dumps(persona, ensure_ascii=False, indent=2) + "\n",
            "strict_rubrics.json": json.dumps(strict, ensure_ascii=False, indent=2) + "\n",
            "original_criteria.json": json.dumps(official, ensure_ascii=False, indent=2) + "\n",
            "preference_units.json": json.dumps(units, ensure_ascii=False, indent=2) + "\n",
            "workbook_annotation.json": json.dumps({
                "high_impact_preferences": record["GT_high_impact_preferences"],
                "average_impact_preferences": record["average_impact_preferences"],
                "acceptable_clarification_paraphrases": record["acceptable_clarification_paraphrases"],
                "simulator_summary": record["full_hidden_persona_for_user_simulator"],
            }, ensure_ascii=False, indent=2) + "\n",
        }
        for name, content in files.items():
            (task_dir / name).write_text(content, encoding="utf-8")
        case = {
            "benchmark_task_id": task_key,
            "pdr_task_id": pdr_task,
            "user_id": user,
            "source_row": row_number,
            "instruction_sha256": hashlib.sha256((task_dir / "instruction.txt").read_bytes()).hexdigest(),
            "strict_rubrics_sha256": hashlib.sha256((task_dir / "strict_rubrics.json").read_bytes()).hexdigest(),
            "official_leaf_count": len(official_rows),
            "askable_high_ids": sorted(ASKABLE[task_key]),
        }
        (case_root / "case.json").write_text(json.dumps(case, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        manifest.append(case)
    manifest.sort(key=lambda item: SELECTED.index(item["benchmark_task_id"]))
    if [item["benchmark_task_id"] for item in manifest] != SELECTED:
        raise RuntimeError("Selected batch is incomplete")
    batch = {
        "schema_version": "0.86",
        "selection_rule": "first five workbook-order released exact pairs excluding previously observed T09",
        "workbook_sha256": hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
        "cases": manifest,
    }
    (ROOT / "batch_manifest.json").write_text(json.dumps(batch, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(batch, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
