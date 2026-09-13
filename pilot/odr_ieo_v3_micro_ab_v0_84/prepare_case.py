#!/usr/bin/env python3
"""Freeze Task 9 / PDR Task 21 / User12 from the user-provided v0.82 workbook."""

from __future__ import annotations

import hashlib
import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
WORKBOOK = Path(
    "/Users/lora/Documents/Codex/2026-09-12/a/outputs/"
    "0913_pdr_bench_15_cold_start_and_micro_rubrics_v0_82.xlsx"
)
PDR = PROJECT / "tmp/pdr-bench/data"


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def main() -> int:
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=False)
    task_sheet = workbook["15 Tasks"]
    rows = list(task_sheet.iter_rows(values_only=True))
    header = list(rows[0])
    record = None
    source_row = None
    for row_number, row in enumerate(rows[1:], 2):
        if row and row[0] and "Benchmark Task 9 | PDR source Task 21 | User12" in str(row[0]):
            record = dict(zip(header, row))
            source_row = row_number
            break
    if record is None:
        raise RuntimeError("Task 9 / PDR Task 21 / User12 was not found")

    rubric_sheet = workbook["Detailed Rubrics"]
    rubric_rows = list(rubric_sheet.iter_rows(values_only=True))
    rubric_header = list(rubric_rows[3])
    selected = [dict(zip(rubric_header, row)) for row in rubric_rows[4:] if row and row[0] == "T09"]
    strict = [row for row in selected if row["rubric_set"] == "P_strict"]
    official_workbook = [row for row in selected if row["rubric_set"] == "P_official"]
    if len(strict) != 67 or len(official_workbook) != 33:
        raise RuntimeError(f"Unexpected rubric counts: strict={len(strict)}, official={len(official_workbook)}")

    official = next(
        item for item in read_jsonl(PDR / "criteria_data/criteria150_en.jsonl")
        if item.get("taskid") == 21 and item.get("userid") == "User12"
    )
    persona = next(
        item for item in read_jsonl(PDR / "persona_data/personas_en.jsonl")
        if item.get("userid") == "User12"
    )

    task_dir = ROOT / "task"
    task_dir.mkdir(parents=True, exist_ok=True)
    instruction = str(record["input/task_instruction for cold start"]).strip()
    (task_dir / "instruction.txt").write_text(instruction + "\n", encoding="utf-8")
    (task_dir / "hidden_persona.json").write_text(
        json.dumps(persona, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (task_dir / "workbook_annotation.json").write_text(
        json.dumps(
            {
                "source_enriched_task_instruction": record[
                    "source_enriched_task_instruction (not cold-start input)"
                ],
                "high_impact_preferences": record["GT_high_impact_preferences"],
                "average_impact_preferences": record["average_impact_preferences"],
                "acceptable_clarification_paraphrases": record[
                    "acceptable_clarification_paraphrases"
                ],
                "simulator_summary": record["full_hidden_persona_for_user_simulator"],
            },
            ensure_ascii=False,
            indent=2,
        )
        + "\n",
        encoding="utf-8",
    )
    (task_dir / "strict_rubrics.json").write_text(
        json.dumps(strict, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (task_dir / "official_workbook_rubrics.json").write_text(
        json.dumps(official_workbook, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (task_dir / "original_criteria.json").write_text(
        json.dumps(official, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    units = [
        {"id": "T9-P1", "impact": "high", "askable_high": True, "label": "moderate risk with explicit downside control"},
        {"id": "T9-P2", "impact": "high", "askable_high": True, "label": "stable long-term compounding and disciplined holding"},
        {"id": "T9-P3", "impact": "high", "askable_high": True, "label": "meaningful technology and innovation exposure"},
        {"id": "T9-P4", "impact": "high", "askable_high": False, "label": "data and financial-analysis-driven decisions", "boundary": "partly an agent/report best practice"},
        {"id": "T9-P5", "impact": "high", "askable_high": False, "label": "diversified portfolio-level risk management", "boundary": "diversification is explicit in the task"},
        {"id": "T9-A1", "impact": "average", "askable_high": False, "label": "maintain reasonable liquidity and reserves"},
        {"id": "T9-A2", "impact": "average", "askable_high": False, "label": "manageable rather than high-frequency monitoring"},
        {"id": "T9-A3", "impact": "average", "askable_high": False, "label": "selective innovative or early-stage exposure"},
    ]
    (task_dir / "preference_units.json").write_text(
        json.dumps(units, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    provenance = {
        "schema_version": "0.84",
        "workbook_path": str(WORKBOOK),
        "workbook_sha256": hashlib.sha256(WORKBOOK.read_bytes()).hexdigest(),
        "source_sheet": "15 Tasks",
        "source_row": source_row,
        "cold_start_column": "K",
        "rubric_sheet": "Detailed Rubrics",
        "strict_leaf_count": len(strict),
        "official_leaf_count": len(official_workbook),
        "instruction_sha256": hashlib.sha256((task_dir / "instruction.txt").read_bytes()).hexdigest(),
        "strict_rubrics_sha256": hashlib.sha256((task_dir / "strict_rubrics.json").read_bytes()).hexdigest(),
    }
    (task_dir / "provenance.json").write_text(
        json.dumps(provenance, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps(provenance, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
