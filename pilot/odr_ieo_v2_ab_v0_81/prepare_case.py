#!/usr/bin/env python3
"""Mechanically freeze the selected workbook row and official PDR assets."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parent
PROJECT = ROOT.parents[1]
WORKBOOK = PROJECT / "0912_pdr_bench_15_reselected_clarification_gt_with_average_preferences.xlsx"
PDR = PROJECT / "tmp/pdr-bench/data"


def read_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def main() -> int:
    workbook = openpyxl.load_workbook(WORKBOOK, read_only=True, data_only=True)
    sheet = workbook["15 Tasks"]
    header = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
    row = next(sheet.iter_rows(min_row=10, max_row=10, values_only=True))
    record = dict(zip(header, row))
    if not str(record["task_instruction"]).startswith("[Benchmark Task 9 | PDR source Task 21 | User12]"):
        raise RuntimeError("selected workbook row changed")

    criteria = next(
        item for item in read_jsonl(PDR / "criteria_data/criteria150_en.jsonl")
        if item.get("taskid") == 21 and item.get("userid") == "User12"
    )
    persona = next(
        item for item in read_jsonl(PDR / "persona_data/personas_en.jsonl")
        if item.get("userid") == "User12"
    )

    task_dir = ROOT / "task"
    task_dir.mkdir(parents=True, exist_ok=True)
    (task_dir / "instruction.txt").write_text(str(record["task_instruction"]).strip() + "\n", encoding="utf-8")
    (task_dir / "hidden_persona.json").write_text(json.dumps(persona, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (task_dir / "original_criteria.json").write_text(json.dumps(criteria, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    (task_dir / "workbook_annotation.json").write_text(
        json.dumps(
            {
                "high_impact_preferences": record["GT_high_impact_preferences"],
                "average_impact_preferences": record["average_impact_preferences"],
                "rubrics": record["rubrics"],
                "acceptable_clarification_paraphrases": record["acceptable_clarification_paraphrases"],
                "simulator_summary": record["full_hidden_persona_for_user_simulator"],
            },
            ensure_ascii=False,
            indent=2,
        ) + "\n",
        encoding="utf-8",
    )
    units = [
        {"id": "T9-P1", "impact": "high", "label": "moderate risk with explicit downside control"},
        {"id": "T9-P2", "impact": "high", "label": "stable long-term compounding and disciplined holding"},
        {"id": "T9-P3", "impact": "high", "label": "meaningful technology and innovation exposure"},
        {"id": "T9-P4", "impact": "high", "label": "data and financial-analysis-driven decisions"},
        {"id": "T9-P5", "impact": "high", "label": "diversified portfolio-level risk management"},
        {"id": "T9-A1", "impact": "average", "label": "maintain reasonable liquidity and reserves"},
        {"id": "T9-A2", "impact": "average", "label": "manageable rather than high-frequency monitoring"},
        {"id": "T9-A3", "impact": "average", "label": "selective innovative or early-stage exposure"},
    ]
    (task_dir / "preference_units.json").write_text(json.dumps(units, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({"criteria_id": criteria["id"], "criteria_leaves": sum(map(len, criteria["personalization_criterions"].values())), "units": len(units)}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
