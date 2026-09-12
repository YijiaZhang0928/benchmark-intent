#!/usr/bin/env python3
"""Extract the first three frozen clarification cases from the source workbook."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from openpyxl import load_workbook


FIELD_NAMES = [
    "task_instruction",
    "GT_high_impact_preferences",
    "full_hidden_persona_for_user_simulator",
    "rubrics",
    "acceptable_clarification_paraphrases",
    "average_impact_preferences",
]


def sha256_text(value: str) -> str:
    return hashlib.sha256(value.encode("utf-8")).hexdigest()


def parse_numbered_lines(text: str, pattern: str) -> list[dict[str, str]]:
    items: list[dict[str, str]] = []
    for line in text.splitlines():
        match = re.match(pattern, line.strip())
        if match:
            items.append({"id": match.group(1), "text": match.group(2).strip()})
    return items


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--output-dir", required=True, type=Path)
    args = parser.parse_args()

    workbook_path = args.workbook.expanduser().resolve()
    output_dir = args.output_dir.expanduser().resolve()
    input_dir = output_dir / "inputs"
    input_dir.mkdir(parents=True, exist_ok=True)

    workbook = load_workbook(workbook_path, read_only=False, data_only=False)
    worksheet = workbook["15 Tasks"]
    headers = [worksheet.cell(1, column).value for column in range(1, 7)]
    if headers != FIELD_NAMES:
        raise ValueError(f"Unexpected source columns: {headers}")

    cases: list[dict] = []
    input_hashes: dict[str, str] = {}
    for task_index, row_index in enumerate(range(2, 5), start=1):
        row = {
            FIELD_NAMES[column_index - 1]: str(worksheet.cell(row_index, column_index).value or "").strip()
            for column_index in range(1, 7)
        }
        preferences = parse_numbered_lines(
            row["GT_high_impact_preferences"], rf"^(T{task_index}-P\d+)\s+(.+)$"
        )
        rubrics = parse_numbered_lines(row["rubrics"], rf"^(T{task_index}-R\d+)\s+(.+)$")
        average_preferences = parse_numbered_lines(
            row["average_impact_preferences"], rf"^(T{task_index}-A\d+)\s+(.+)$"
        )
        if len(preferences) != 5 or len(rubrics) != 5:
            raise ValueError(
                f"Task {task_index}: expected five high-impact preferences and five rubrics, "
                f"found {len(preferences)} and {len(rubrics)}"
            )

        task_id = f"T{task_index}"
        instruction_path = input_dir / f"{task_id}_instruction.txt"
        full_persona_path = input_dir / f"{task_id}_full_persona.txt"
        instruction_path.write_text(row["task_instruction"] + "\n", encoding="utf-8")
        full_persona_text = (
            row["task_instruction"]
            + "\n\nUSER CONTEXT\n\n"
            + row["full_hidden_persona_for_user_simulator"]
            + "\n"
        )
        full_persona_path.write_text(full_persona_text, encoding="utf-8")
        input_hashes[str(instruction_path.relative_to(output_dir))] = sha256_text(
            row["task_instruction"] + "\n"
        )
        input_hashes[str(full_persona_path.relative_to(output_dir))] = sha256_text(full_persona_text)

        cases.append(
            {
                "case_id": task_id,
                "source_sheet": "15 Tasks",
                "source_row": row_index,
                "source_cells": {
                    "task_instruction": f"A{row_index}",
                    "high_impact_preferences": f"B{row_index}",
                    "hidden_persona": f"C{row_index}",
                    "rubrics": f"D{row_index}",
                    "acceptable_paraphrases": f"E{row_index}",
                    "average_impact_preferences": f"F{row_index}",
                },
                "instruction": row["task_instruction"],
                "high_impact_preferences": preferences,
                "hidden_persona": row["full_hidden_persona_for_user_simulator"],
                "rubrics": rubrics,
                "acceptable_clarification_paraphrases": row[
                    "acceptable_clarification_paraphrases"
                ],
                "average_impact_preferences": average_preferences,
                "instruction_sha256": sha256_text(row["task_instruction"]),
            }
        )

    payload = {
        "schema_version": "0.76",
        "source_workbook_name": workbook_path.name,
        "source_workbook_sha256": hashlib.sha256(workbook_path.read_bytes()).hexdigest(),
        "source_sheet": "15 Tasks",
        "source_rows": [2, 3, 4],
        "cases": cases,
        "generated_input_hashes": input_hashes,
    }
    (output_dir / "cases.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(json.dumps({"cases": len(cases), "output": str(output_dir / "cases.json")}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
